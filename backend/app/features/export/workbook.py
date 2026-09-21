"""Renders the literature workbook (.xlsx) from an export snapshot.

Formatting matches the workbook this app already produced: frozen header row,
bold white-on-indigo headers, sized columns, wrapped body text.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.features.export import snapshot as snapshot_mod

HEADERS = [
    "Paper Title",
    "Citation (Author, Year)",
    "Research Objective / Questions",
    "Methodology & Sample",
    "Key Findings",
    "Limitations & Gaps",
    "Relevance / Contribution",
    "APA Reference",
    "Authors",
    "Year",
    "DOI",
    "Journal",
    "Missing metadata",
    "Source",
]

WIDTHS = [42, 22, 46, 42, 46, 42, 42, 52, 28, 8, 24, 26, 24, 16]


def _row(paper) -> list:
    entry = paper.entry.values if paper.entry else {}
    return [
        paper.title,
        paper.citation,
        entry.get("research_objective", snapshot_mod.NOT_GENERATED),
        entry.get("methodology", snapshot_mod.NOT_GENERATED),
        entry.get("key_findings", snapshot_mod.NOT_GENERATED),
        entry.get("limitations", snapshot_mod.NOT_GENERATED),
        entry.get("relevance", snapshot_mod.NOT_GENERATED),
        paper.apa_reference,
        "; ".join(paper.authors) if paper.authors else snapshot_mod.AUTHOR_UNKNOWN,
        # Year/DOI/Journal stay blank when absent so the columns remain usable
        # as data; the "Missing metadata" and "Source" columns carry why.
        paper.year,
        paper.doi or "",
        paper.journal or "",
        ", ".join(paper.missing_fields) if paper.missing_fields else "",
        paper.provenance,
    ]


def render(snapshot) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Literature Mapping"
    sheet.append(HEADERS)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")
    for col, width in enumerate(WIDTHS, start=1):
        cell = sheet.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "A2"

    for paper in snapshot.papers:
        sheet.append(_row(paper))

    for row_index in range(2, sheet.max_row + 1):
        for col_index in range(1, len(HEADERS) + 1):
            sheet.cell(row=row_index, column=col_index).alignment = Alignment(
                wrap_text=True, vertical="top"
            )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
