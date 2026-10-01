"""Renders a saved chat session as a take-away report.

The session report is the second artifact of the export sprint: the questions
the researcher asked, the answers NARA gave, and — under each answer — the
sources it rested on, with the citation labels captured at answer time.

Citations to a document that has since been deleted are the point of this
module. ``documents`` deletion never touches ``chat_messages.sources``, so a
saved answer can still cite a paper that is gone. Nothing is silently dropped:
such a source is marked ``not in this project any more`` and counted on the
cover, so the researcher sees both what the answer rested on and that the
underlying paper has left her library.

Reads SQLite directly rather than importing the chat feature, which honours the
rule that features never import from each other.
"""

import json
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from typing import Iterable, Optional

from docx import Document

from app.features.export.snapshot import NOT_IN_PROJECT

_UNTITLED = "Untitled conversation"


@dataclass
class SourceRef:
    chunk_id: str
    title: str
    page: Optional[int]
    apa_reference: str
    document_id: Optional[str]
    present: bool

    @property
    def label(self) -> str:
        """The citation as it was captured, with the page when there was one."""
        text = self.apa_reference or self.title or self.chunk_id
        if self.page is not None:
            text = f"{text} (p. {self.page})"
        if not self.present:
            text = f"{text} — {NOT_IN_PROJECT}"
        return text


@dataclass
class Message:
    role: str
    content: str
    model_used: Optional[str]
    created_at: str
    sources: list[SourceRef] = field(default_factory=list)


@dataclass
class Conversation:
    session_id: str
    title: str
    project_name: str
    model_used: Optional[str]
    generated_at: str
    messages: list[Message]
    unresolved: list[str] = field(default_factory=list)


def _parse_sources(raw: Optional[str]) -> list[dict]:
    if not raw:
        return []
    try:
        parsed = json.loads(raw)
    except (TypeError, ValueError):
        return []
    return [item for item in parsed if isinstance(item, dict)] if isinstance(parsed, list) else []


def _existing_document_ids(db, user_id: str, document_ids: Iterable[str]) -> set[str]:
    """The subset of ``document_ids`` that still exist, in batched queries.

    A batch rather than a SELECT per source, so a session with many citations is
    one round-trip per chunk instead of one per source. A source with no document
    id (a web fallback) never had a project document, so it is not a
    deleted-document case and is reported as present.
    """
    ids = {document_id for document_id in document_ids if document_id}
    if not ids:
        return set()
    values = list(ids)
    found: set[str] = set()
    chunk = 500  # stay well under SQLite's bound-variable limit
    for start in range(0, len(values), chunk):
        batch = values[start : start + chunk]
        placeholders = ", ".join("?" for _ in batch)
        rows = db.execute(
            f"SELECT id FROM documents WHERE user_id = ? AND id IN ({placeholders})",
            (user_id, *batch),
        ).fetchall()
        found.update(row["id"] for row in rows)
    return found


def _build_source(item: dict, present_ids: set[str]) -> SourceRef:
    document_id = item.get("document_id")
    return SourceRef(
        chunk_id=str(item.get("chunk_id") or ""),
        title=str(item.get("title") or ""),
        page=item.get("page"),
        apa_reference=str(item.get("apa_reference") or ""),
        document_id=document_id,
        present=True if not document_id else document_id in present_ids,
    )


def build(db, user_id: str, session_id: str) -> Optional[Conversation]:
    """Reads one saved session into a report snapshot. Never touches the network."""
    session = db.execute(
        "SELECT * FROM chat_sessions WHERE id = ? AND user_id = ?", (session_id, user_id)
    ).fetchone()
    if session is None:
        return None
    session = dict(session)

    project_name = "Project"
    if session.get("project_id"):
        project = db.execute(
            "SELECT name FROM projects WHERE id = ? AND user_id = ?",
            (session["project_id"], user_id),
        ).fetchone()
        if project is not None:
            project_name = ((project["name"] or "") or "Project").strip()

    rows = db.execute(
        """SELECT * FROM chat_messages
           WHERE session_id = ? AND role IN ('user', 'assistant')
           ORDER BY created_at, rowid""",
        (session_id,),
    ).fetchall()

    parsed = []
    for row in rows:
        message_row = dict(row)
        parsed.append((message_row, _parse_sources(message_row.get("sources"))))
    cited_ids = {
        item.get("document_id")
        for _, sources in parsed
        for item in sources
        if item.get("document_id")
    }
    present_ids = _existing_document_ids(db, user_id, cited_ids)

    messages = [
        Message(
            role=row.get("role") or "",
            content=(row.get("content") or "").strip(),
            model_used=row.get("model_used"),
            created_at=row.get("created_at") or "",
            sources=[_build_source(item, present_ids) for item in sources],
        )
        for row, sources in parsed
    ]

    missing = sum(1 for message in messages for source in message.sources if not source.present)
    unresolved = []
    if not messages:
        unresolved.append("This conversation has no messages yet.")
    if missing:
        unresolved.append(
            f"{missing} citation(s) point at documents no longer in this project."
        )

    return Conversation(
        session_id=session_id,
        title=(session.get("title") or "").strip() or _UNTITLED,
        project_name=project_name,
        model_used=session.get("model_used"),
        generated_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
        messages=messages,
        unresolved=unresolved,
    )


def _add_cover(document, conversation: Conversation) -> None:
    document.add_heading(conversation.title, level=0)
    info = document.add_paragraph()
    info.add_run(
        f"{conversation.project_name} · Exported {conversation.generated_at} · "
        f"{len(conversation.messages)} message(s)"
    ).italic = True
    if conversation.model_used:
        info.add_run(f" · model: {conversation.model_used}").italic = True

    document.add_heading("Before you rely on this", level=1)
    if conversation.unresolved:
        document.add_paragraph(
            "The app could not fully resolve the following. They are shown below "
            "rather than dropped or filled in with a guess:"
        )
        for item in conversation.unresolved:
            document.add_paragraph(item, style="List Bullet")
    else:
        document.add_paragraph("Nothing was missing when this export ran.")


def _add_message(document, message: Message) -> None:
    speaker = "You" if message.role == "user" else "NARA"
    document.add_heading(speaker, level=2)

    lines = message.content.splitlines() or [""]
    for line in lines:
        document.add_paragraph(line)

    if message.sources:
        note = document.add_paragraph()
        note.add_run("Sources:").italic = True
        for source in message.sources:
            paragraph = document.add_paragraph(source.label, style="List Bullet")
            if not source.present:
                for run in paragraph.runs:
                    run.italic = True


def render_docx(conversation: Conversation) -> bytes:
    document = Document()
    _add_cover(document, conversation)
    for message in conversation.messages:
        _add_message(document, message)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
