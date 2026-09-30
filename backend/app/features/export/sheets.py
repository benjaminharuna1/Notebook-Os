"""Shared openpyxl sheet formatting.

One definition of what an export sheet looks like (header band, frozen top row,
sized columns, wrapped body) so every sheet in every workbook matches.
"""

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")


def write_sheet(sheet, headers: list, widths: list, rows: list) -> None:
    sheet.append(headers)
    for col, width in enumerate(widths, start=1):
        cell = sheet.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "A2"

    for row in rows:
        sheet.append(row)

    for row_index in range(2, sheet.max_row + 1):
        for col_index in range(1, len(headers) + 1):
            sheet.cell(row=row_index, column=col_index).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
