import json
import math
import re
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

from rapidfuzz import fuzz

from app.features.embedding.providers.factory import resolve_embedding_provider
from app.features.graph.builder import layout_nodes
from app.features.literature import metadata as metadata_sources
from app.features.literature.metadata import (
    MAX_CANDIDATES,
    extract_doi,
    extract_year,
    heuristic_title,
)
from app.features.literature.schemas import ClusterInfo, LiteratureMapResponse, PaperEdge, PaperNode
from app.features.settings.service import SettingsService
from app.shared.id_utils import generate_id

SIMILARITY_THRESHOLD = 0.5
SIMILARITY_TOP_K = 5
CITATION_THRESHOLD = 85
VERIFIED_TITLE_THRESHOLD = 90
MAX_REFERENCE_LINES = 200

_CREATE_REFERENCES_TABLE = """
CREATE TABLE IF NOT EXISTS paper_references (
    id               TEXT PRIMARY KEY,
    paper_id         TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
    raw_ref          TEXT NOT NULL,
    matched_paper_id TEXT,
    confidence       REAL
)
"""
_ALTER_MAP_TYPE = "ALTER TABLE graph_history ADD COLUMN map_type TEXT DEFAULT 'concepts'"

_CREATE_ENTRIES_TABLE = """
CREATE TABLE IF NOT EXISTS literature_entries (
    paper_id           TEXT PRIMARY KEY REFERENCES documents(id) ON DELETE CASCADE,
    user_id            TEXT NOT NULL,
    project_id         TEXT NOT NULL,
    citation           TEXT,
    research_objective TEXT,
    methodology        TEXT,
    key_findings       TEXT,
    limitations        TEXT,
    relevance          TEXT,
    apa_reference      TEXT,
    auto_generated     INTEGER NOT NULL DEFAULT 0,
    user_edited        TEXT,
    updated_at         DATETIME DEFAULT CURRENT_TIMESTAMP
)
"""

_ENTRY_FIELDS = (
    "citation",
    "research_objective",
    "methodology",
    "key_findings",
    "limitations",
    "relevance",
    "apa_reference",
)


class LiteratureService:
    """Builds a paper-level literature map for a project.

    Papers (the project's indexed documents) become nodes. Two typed edge kinds
    connect them:
    * ``similarity`` — cosine similarity of the papers' embeddings (title +
      abstract via the user's active embedding provider).
    * ``citation`` — a paper's reference list fuzzy-matched against the other
      papers' titles.

    Nodes are grouped into communities (greedy modularity) and laid out with
    the same force-directed pass the concept graph uses. Metadata is enriched
    from Crossref (best-effort; failures degrade to ``unverified``).
    """

    def __init__(self, db):
        self.db = db
        cur = self.db.cursor()
        cur.execute(_CREATE_REFERENCES_TABLE)
        cur.execute(_CREATE_ENTRIES_TABLE)
        try:
            cur.execute(_ALTER_MAP_TYPE)
        except Exception:
            pass
        try:
            cur.execute("ALTER TABLE literature_entries ADD COLUMN user_edited TEXT")
        except Exception:
            pass
        try:
            cur.execute(
                "ALTER TABLE documents ADD COLUMN metadata_user_edited INTEGER DEFAULT 0"
            )
        except Exception:
            pass
        try:
            cur.execute("ALTER TABLE documents ADD COLUMN extracted_doi TEXT")
        except Exception:
            pass
        try:
            cur.execute("ALTER TABLE documents ADD COLUMN metadata_candidates TEXT")
        except Exception:
            pass
        for column in ("pages", "journal", "volume", "issue", "publisher", "url"):
            try:
                cur.execute(f"ALTER TABLE documents ADD COLUMN {column} TEXT")
            except Exception:
                pass
        self.db.commit()

    # --- papers -------------------------------------------------------------

    def papers(self, user_id: str, project_id: str) -> List[dict]:
        rows = self.db.execute(
            """SELECT id, title, author, year, doi, abstract, verification_status,
                      apa_reference, authors, metadata_user_edited, filename,
                      journal, volume, issue, pages, publisher, url
               FROM documents
               WHERE user_id = ? AND project_id = ?
               ORDER BY title COLLATE NOCASE""",
            (user_id, project_id),
        ).fetchall()
        out = []
        for row in rows:
            out.append(self._paper_from_row(row))
        return out

    @staticmethod
    def _paper_from_row(row) -> dict:
        stored_authors = row["authors"]
        try:
            authors = json.loads(stored_authors) if stored_authors else []
        except (TypeError, ValueError):
            authors = []
        if not authors and row["author"]:
            authors = [a.strip() for a in re.split(r"[;,]", row["author"]) if a.strip()]
        return {
            "id": row["id"],
            "title": row["title"] or "",
            "author": row["author"],
            "year": row["year"],
            "doi": row["doi"],
            "abstract": row["abstract"],
            "verification_status": row["verification_status"],
            "apa_reference": row["apa_reference"],
            "authors": authors,
            "metadata_user_edited": bool(row["metadata_user_edited"]),
            "filename": row["filename"] if "filename" in row.keys() else None,
            "journal": row["journal"] if "journal" in row.keys() else None,
            "volume": row["volume"] if "volume" in row.keys() else None,
            "issue": row["issue"] if "issue" in row.keys() else None,
            "pages": row["pages"] if "pages" in row.keys() else None,
            "publisher": row["publisher"] if "publisher" in row.keys() else None,
            "url": row["url"] if "url" in row.keys() else None,
        }

    # --- editable literature entries ----------------------------------------

    def entries(self, user_id: str, project_id: str) -> List[dict]:
        rows = self.db.execute(
            """SELECT e.*, d.title AS title
               FROM literature_entries e
               JOIN documents d ON d.id = e.paper_id
               WHERE e.user_id = ? AND e.project_id = ?
               ORDER BY e.citation COLLATE NOCASE, d.title COLLATE NOCASE""",
            (user_id, project_id),
        ).fetchall()
        return [self._entry_from_row(r) for r in rows]

    def get_entry(self, paper_id: str, user_id: str) -> Optional[dict]:
        row = self.db.execute(
            """SELECT e.*, d.title AS title
               FROM literature_entries e
               JOIN documents d ON d.id = e.paper_id
               WHERE e.paper_id = ? AND e.user_id = ?""",
            (paper_id, user_id),
        ).fetchone()
        return self._entry_from_row(row) if row else None

    @staticmethod
    def _entry_from_row(row) -> dict:
        return {
            "paper_id": row["paper_id"],
            "title": row["title"] or "",
            "citation": row["citation"],
            "research_objective": row["research_objective"],
            "methodology": row["methodology"],
            "key_findings": row["key_findings"],
            "limitations": row["limitations"],
            "relevance": row["relevance"],
            "apa_reference": row["apa_reference"],
            "auto_generated": bool(row["auto_generated"]),
        }

    def upsert_entry(
        self,
        paper_id: str,
        user_id: str,
        project_id: str,
        fields: dict,
        auto: bool = False,
        auto_generated: Optional[bool] = None,
        overwrite: Optional[set] = None,
    ) -> Optional[dict]:
        """Creates or updates a paper's editable literature entry.

        In ``auto`` mode only empty fields are filled in, so user edits are
        never clobbered by a rebuild. In user mode the provided values replace
        the stored ones (empty strings clear a field). ``auto_generated`` forces
        the badge flag when a caller wants to replace content with fresh AI
        output without marking it as a manual edit. ``overwrite`` names fields
        that are written even in ``auto`` mode (used when source metadata
        changes); those fields are not counted as user edits.
        """
        overwrite = overwrite or set()
        clean = {}
        for key in _ENTRY_FIELDS:
            value = fields.get(key)
            clean[key] = value.strip() if isinstance(value, str) else value

        existing = self.db.execute(
            "SELECT * FROM literature_entries WHERE paper_id = ? AND user_id = ?",
            (paper_id, user_id),
        ).fetchone()

        user_edited = self._parse_user_edited(existing["user_edited"] if existing else None)

        if auto:
            merged = {}
            for key in _ENTRY_FIELDS:
                value = clean.get(key)
                if value is None:
                    continue
                if key in overwrite:
                    merged[key] = value
                    continue
                current = ((existing[key] or "").strip() if existing else "").strip()
                if not current:
                    merged[key] = value
        else:
            merged = {k: v for k, v in clean.items() if v is not None}
            user_edited = user_edited | set(merged.keys())

        if existing:
            if merged:
                sets = ", ".join(f"{k} = ?" for k in merged)
                flag = existing["auto_generated"] if auto else 0
                if auto_generated is not None:
                    flag = 1 if auto_generated else 0
                self.db.execute(
                    f"UPDATE literature_entries SET {sets}, auto_generated = ?, "
                    "user_edited = ?, updated_at = CURRENT_TIMESTAMP WHERE paper_id = ?",
                    (*merged.values(), flag, json.dumps(sorted(user_edited)), paper_id),
                )
                self.db.commit()
        else:
            row = {k: None for k in _ENTRY_FIELDS}
            for key, value in merged.items():
                row[key] = value
            self.db.execute(
                """INSERT INTO literature_entries
                      (paper_id, user_id, project_id, citation, research_objective,
                       methodology, key_findings, limitations, relevance, apa_reference,
                       auto_generated, user_edited)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    paper_id,
                    user_id,
                    project_id,
                    row["citation"],
                    row["research_objective"],
                    row["methodology"],
                    row["key_findings"],
                    row["limitations"],
                    row["relevance"],
                    row["apa_reference"],
                    1 if (auto or auto_generated) else 0,
                    json.dumps(sorted(user_edited)) if user_edited else None,
                ),
            )
            self.db.commit()
        return self.get_entry(paper_id, user_id)

    @staticmethod
    def _parse_user_edited(raw: Optional[str]) -> set:
        if not raw:
            return set()
        try:
            value = json.loads(raw)
            return set(value) if isinstance(value, list) else set()
        except (TypeError, ValueError):
            return set()

    @staticmethod
    def auto_citation(paper: dict) -> str:
        """Derives a `Author, Year` citation from the paper's metadata."""
        authors = paper.get("authors") or []
        family = None
        if authors:
            first = authors[0]
            family = first.split(",", 1)[0].strip() if ", " in first else (first.split()[-1] or None)
        if not family and paper.get("author"):
            parts = [a for a in re.split(r"[;,]", paper["author"]) if a.strip()]
            if parts:
                family = parts[0].split()[-1]
        if not family:
            family = (paper.get("title") or "Untitled")[:30]
        year = paper.get("year")
        return f"{family}, {year}" if year else f"{family}, n.d."

    def ensure_entries(self, user_id: str, project_id: str) -> int:
        """Seeds a literature_entries row per paper with citation + APA filled.

        Runs before LLM summarization so the table always has stable, editable
        identity columns for every paper.
        """
        papers = self.papers(user_id, project_id)
        for paper in papers:
            self.upsert_entry(
                paper["id"],
                user_id,
                project_id,
                {
                    "citation": self.auto_citation(paper),
                    "apa_reference": paper.get("apa_reference") or self.apa_reference(paper),
                },
                auto=True,
            )
        return len(papers)

    def regenerate_metadata(
        self, user_id: str, project_id: str, paper_ids: Optional[List[str]] = None
    ) -> List[dict]:
        """Re-runs enrichment (DOI/title/AI lookup) for selected papers, or all
        papers when ``paper_ids`` is None, and refreshes their literature
        entries' citation/APA. Returns a per-paper result summary.

        User-edited papers are skipped by ``enrich`` itself; everything here is
        best-effort and never raises per paper.
        """
        papers = self.papers(user_id, project_id)
        if paper_ids:
            selected = {pid for pid in paper_ids}
            papers = [p for p in papers if p["id"] in selected]
        results = []
        for paper in papers:
            try:
                self.enrich(paper, user_id=user_id)
                self.save_enrichment(paper)
                self.upsert_entry(
                    paper["id"],
                    user_id,
                    project_id,
                    {
                        "citation": self.auto_citation(paper),
                        "apa_reference": paper.get("apa_reference")
                        or self.apa_reference(paper),
                    },
                    auto=True,
                )
                results.append(
                    {
                        "paper_id": paper["id"],
                        "status": paper.get("verification_status"),
                    }
                )
            except Exception:
                results.append({"paper_id": paper["id"], "status": "error"})
        return results

    # --- editable paper metadata --------------------------------------------

    def get_metadata(self, paper_id: str, user_id: str, project_id: str) -> Optional[dict]:
        """Returns the raw paper metadata that feeds the APA reference, plus
        any unresolved metadata candidates for the paper."""
        row = self.db.execute(
            """SELECT id, title, author, year, doi, abstract, verification_status,
                      apa_reference, authors, metadata_user_edited, file_type,
                      extracted_doi, metadata_candidates,
                      journal, volume, issue, pages, publisher, url
               FROM documents
               WHERE id = ? AND user_id = ? AND project_id = ?""",
            (paper_id, user_id, project_id),
        ).fetchone()
        if not row:
            return None
        return self._metadata_from_row(row)

    @staticmethod
    def _parse_candidates(raw) -> List[dict]:
        if not raw:
            return []
        try:
            value = json.loads(raw)
            return value if isinstance(value, list) else []
        except (TypeError, ValueError):
            return []

    @staticmethod
    def _metadata_from_row(row) -> dict:
        paper = LiteratureService._paper_from_row(row)
        candidates = LiteratureService._parse_candidates(row["metadata_candidates"])
        return {
            "title": paper["title"],
            "authors": paper["authors"],
            "year": paper["year"],
            "doi": paper["doi"],
            "abstract": paper["abstract"],
            "journal": paper["journal"],
            "volume": paper["volume"],
            "issue": paper["issue"],
            "pages": paper["pages"],
            "publisher": paper["publisher"],
            "url": paper["url"],
            "apa_reference": paper["apa_reference"] or LiteratureService.apa_reference(paper),
            "verification_status": paper["verification_status"],
            "metadata_user_edited": paper["metadata_user_edited"],
            "file_type": row["file_type"] if "file_type" in row.keys() else None,
            "extracted_doi": row["extracted_doi"] if "extracted_doi" in row.keys() else None,
            "candidates": candidates,
        }

    def apply_candidate(
        self, paper_id: str, user_id: str, project_id: str, index: int
    ) -> Optional[dict]:
        """Adopts a stored metadata candidate (chosen by the user) as the
        paper's verified metadata. Freezes it so rebuilds keep the choice."""
        meta = self.get_metadata(paper_id, user_id, project_id)
        if meta is None or not (0 <= index < len(meta.get("candidates") or [])):
            return None
        record = meta["candidates"][index]
        fields = {
            "title": record.get("title"),
            "authors": record.get("authors") or [],
            "year": record.get("year"),
            "doi": record.get("doi"),
            "abstract": record.get("abstract"),
            "journal": record.get("journal") or record.get("container_title"),
            "volume": record.get("volume"),
            "issue": record.get("issue"),
            "pages": record.get("pages"),
            "publisher": record.get("publisher"),
            "url": record.get("url"),
        }
        updated = self.update_metadata(paper_id, user_id, project_id, fields)
        if updated is None:
            return None
        self.db.execute(
            """UPDATE documents
               SET verification_status = 'verified', metadata_candidates = '[]'
               WHERE id = ? AND user_id = ? AND project_id = ?""",
            (paper_id, user_id, project_id),
        )
        self.db.commit()
        return self.get_metadata(paper_id, user_id, project_id)

    def update_metadata(
        self, paper_id: str, user_id: str, project_id: str, fields: dict
    ) -> Optional[dict]:
        """Applies user edits to a paper's metadata and regenerates the APA.

        The document's author/year/DOI/abstract/title are stored as given, the
        APA reference is recomputed from them, and the literature entry's
        ``citation``/``apa_reference`` refresh too — unless the user explicitly
        overrode those two cells in the literature table.
        """
        paper = self.get_metadata(paper_id, user_id, project_id)
        if paper is None:
            return None
        for key in ("authors", "year", "doi", "abstract", "title", "journal",
                    "volume", "issue", "pages", "publisher", "url"):
            if key not in fields:
                continue
            value = fields[key]
            if key == "authors":
                cleaned = [a.strip() for a in value if isinstance(a, str) and a.strip()]
                paper["authors"] = cleaned
            elif key == "title":
                paper["title"] = (value or "").strip() or paper["title"]
            elif key == "year":
                paper["year"] = value
            else:
                paper[key] = (value or "").strip() or None

        apa = self.apa_reference(
            {
                "title": paper["title"],
                "author": None,
                "year": paper["year"],
                "doi": paper["doi"],
                "authors": paper["authors"],
                "journal": paper.get("journal"),
                "volume": paper.get("volume"),
                "issue": paper.get("issue"),
                "pages": paper.get("pages"),
                "publisher": paper.get("publisher"),
                "url": paper.get("url"),
            }
        )
        self.db.execute(
            """UPDATE documents
               SET title = ?, authors = ?, year = ?, doi = ?, abstract = ?,
                   journal = ?, volume = ?, issue = ?, pages = ?, publisher = ?,
                   url = ?, apa_reference = ?, verification_status = NULL,
                   metadata_user_edited = 1
               WHERE id = ?""",
            (
                paper["title"],
                json.dumps(paper["authors"]),
                paper["year"],
                paper["doi"],
                paper["abstract"],
                paper.get("journal"),
                paper.get("volume"),
                paper.get("issue"),
                paper.get("pages"),
                paper.get("publisher"),
                paper.get("url"),
                apa,
                paper_id,
            ),
        )
        self.db.commit()

        raw = self.db.execute(
            "SELECT user_edited FROM literature_entries WHERE paper_id = ? AND user_id = ?",
            (paper_id, user_id),
        ).fetchone()
        user_edited = self._parse_user_edited(raw["user_edited"] if raw else None)
        entry_fields = {}
        if "citation" not in user_edited:
            entry_fields["citation"] = self.auto_citation(paper)
        if "apa_reference" not in user_edited:
            entry_fields["apa_reference"] = apa
        if entry_fields:
            self.upsert_entry(
                paper_id,
                user_id,
                project_id,
                entry_fields,
                auto=True,
                overwrite=set(entry_fields.keys()),
            )
        return self.get_metadata(paper_id, user_id, project_id)

    def export_rows(self, user_id: str, project_id: str) -> List[dict]:
        """One row per paper: title + the seven literature-mapping columns.

        Filled from the editable entries where available, falling back to
        auto-derived citation/APA so the sheet never has blank identity cells.
        """
        papers = self.papers(user_id, project_id)
        entries = {e["paper_id"]: e for e in self.entries(user_id, project_id)}
        rows = []
        for paper in papers:
            entry = entries.get(paper["id"]) or {}
            rows.append(
                {
                    "title": paper.get("title") or "",
                    "citation": (entry.get("citation") or "").strip()
                    or self.auto_citation(paper),
                    "research_objective": (entry.get("research_objective") or "").strip(),
                    "methodology": (entry.get("methodology") or "").strip(),
                    "key_findings": (entry.get("key_findings") or "").strip(),
                    "limitations": (entry.get("limitations") or "").strip(),
                    "relevance": (entry.get("relevance") or "").strip(),
                    "apa_reference": (entry.get("apa_reference") or "").strip()
                    or (paper.get("apa_reference") or "").strip(),
                }
            )
        return rows

    def export_workbook(self, user_id: str, project_id: str) -> bytes:
        from io import BytesIO

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        headers = [
            "Paper Title",
            "Citation (Author, Year)",
            "Research Objective / Questions",
            "Methodology & Sample",
            "Key Findings",
            "Limitations & Gaps",
            "Relevance / Contribution",
            "APA Reference",
        ]
        widths = [42, 22, 46, 42, 46, 42, 42, 52]

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Literature Mapping"
        sheet.append(headers)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F46E5")
        for col, width in enumerate(widths, start=1):
            cell = sheet.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "A2"

        for row in self.export_rows(user_id, project_id):
            sheet.append(
                [
                    row["title"],
                    row["citation"],
                    row["research_objective"],
                    row["methodology"],
                    row["key_findings"],
                    row["limitations"],
                    row["relevance"],
                    row["apa_reference"],
                ]
            )
        for row_index in range(2, sheet.max_row + 1):
            for col_index in range(1, len(headers) + 1):
                sheet.cell(row=row_index, column=col_index).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def export_references_docx(self, user_id: str, project_id: str, project_name: str) -> bytes:
        """Builds a Word document of APA references for every paper.

        Layout: the project name as the document title, a ``References``
        sub-heading, then each reference as a hanging-indent numbered paragraph,
        sorted alphabetically and de-duplicated.
        """
        from io import BytesIO

        from docx import Document
        from docx.shared import Pt

        papers = self.papers(user_id, project_id)
        entries = {e["paper_id"]: e for e in self.entries(user_id, project_id)}
        refs = []
        for paper in papers:
            entry = entries.get(paper["id"]) or {}
            ref = (
                (entry.get("apa_reference") or "").strip()
                or (paper.get("apa_reference") or "").strip()
                or self.apa_reference(paper)
            )
            if ref:
                refs.append(ref)
        refs = sorted({ref for ref in refs if ref})

        document = Document()
        document.add_heading(project_name or "References", level=0)
        document.add_heading("References", level=1)
        if not refs:
            document.add_paragraph("No papers with references yet.")
        for ref in refs:
            paragraph = document.add_paragraph(ref)
            try:
                paragraph.style = document.styles["List Number"]
            except KeyError:
                pass
            paragraph.paragraph_format.left_indent = Pt(36)
            paragraph.paragraph_format.first_line_indent = Pt(-36)
            paragraph.paragraph_format.space_after = Pt(6)

        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    # --- Crossref enrichment ------------------------------------------------

    def enrich(self, paper: dict, user_id: Optional[str] = None) -> dict:
        """Best-effort metadata enrichment. Never raises.

        Order of confidence:
        1. A DOI extracted from the PDF's first page → exact Crossref lookup
           (authoritative, always ``verified``).
        2. A DOI produced by the LLM's first-page reading → exact Crossref lookup.
        3. Title-based search across Crossref + OpenAlex, cross-checked with
           author/year sanity before marking ``verified``.
        4. LLM-extracted metadata (title/authors/year/journal) as a best-effort
           fill-in marked ``ai`` when nothing external corroborates it — the user
           can adjust it in the metadata modal later.
        5. No trustworthy match → ``unverified`` with candidates stored for the
           user to pick from manually.

        Papers whose metadata the user already edited are left untouched.
        """
        if paper.get("metadata_user_edited"):
            return paper

        title = (paper.get("title") or "").strip()
        first_page = self._first_page_text(paper["id"])
        paper["extracted_doi"] = extract_doi(first_page)
        heuristic = heuristic_title(first_page) or title or self._title_from_filename(paper)

        if not heuristic or len(heuristic) < 10:
            paper["verification_status"] = None
            paper["metadata_candidates"] = []
            return paper

        if paper["extracted_doi"]:
            record = metadata_sources.crossref_by_doi(paper["extracted_doi"])
            if record:
                record["title"] = record["title"] or heuristic
                self._apply_record(paper, record, verified=True)
                paper["metadata_candidates"] = []
                return paper

        llm_fields = self._ai_extract(user_id, first_page, paper.get("filename"))
        ai_candidate = self._ai_candidate(llm_fields)
        if ai_candidate and ai_candidate.get("doi"):
            record = metadata_sources.crossref_by_doi(ai_candidate["doi"])
            if record:
                record["title"] = record["title"] or ai_candidate["title"]
                self._apply_record(paper, record, verified=True)
                paper["metadata_candidates"] = []
                return paper

        query_title = (ai_candidate or {}).get("title") or heuristic
        external = metadata_sources.crossref_by_title(query_title)
        external += metadata_sources.openalex_by_title(query_title)

        best, score = self._pick_best(query_title, external)
        if best is not None and score >= CITATION_THRESHOLD:
            verified = score >= VERIFIED_TITLE_THRESHOLD and self._sanity_ok(
                paper, first_page, best, llm_fields
            )
            self._apply_record(paper, best, verified=verified)
            if verified:
                paper["metadata_candidates"] = []
                return paper

        if ai_candidate:
            self._apply_record(paper, ai_candidate, verified=False, fill=True)
            paper["verification_status"] = "ai"
        else:
            paper["verification_status"] = "unverified"
        paper["metadata_candidates"] = (external + ([ai_candidate] if ai_candidate else []))[:MAX_CANDIDATES]
        return paper

    def _ai_extract(
        self, user_id: Optional[str], first_page: Optional[str], filename: Optional[str]
    ) -> dict:
        if not user_id or not first_page:
            return {}
        from app.features.literature.llm_service import LiteratureLLMService

        try:
            llm = LiteratureLLMService(self.db)
            fields = llm.extract_paper_metadata(user_id, first_page, filename or "")
            return fields if isinstance(fields, dict) else {}
        except Exception:
            return {}

    @staticmethod
    def _ai_candidate(fields: dict) -> Optional[dict]:
        title = (fields.get("title") or "").strip()
        if not title or len(title.split()) < 3:
            return None
        authors = fields.get("authors") or []
        if not isinstance(authors, list):
            authors = []
        authors = [a.strip() for a in authors if isinstance(a, str) and a.strip()]
        doi = (fields.get("doi") or "").strip() or None
        if doi:
            doi = (
                doi.replace("https://doi.org/", "")
                .replace("http://doi.org/", "")
                .replace("http://dx.doi.org/", "")
                .rstrip(".,")
                .lower()
                or None
            )
        year = fields.get("year")
        try:
            year = int(year) if year not in (None, "") else None
        except (TypeError, ValueError):
            year = None
        journal = (fields.get("journal") or "").strip() or None
        abstract = (fields.get("abstract") or "").strip() or None
        return {
            "source": "ai",
            "doi": doi,
            "title": title,
            "authors": authors,
            "year": year,
            "abstract": abstract,
            "journal": journal,
            "volume": (fields.get("volume") or "").strip() or None,
            "issue": (fields.get("issue") or "").strip() or None,
            "pages": (fields.get("pages") or "").strip() or None,
            "publisher": (fields.get("publisher") or "").strip() or None,
            "url": (fields.get("url") or "").strip() or None,
        }

    @staticmethod
    def _title_from_filename(paper: dict) -> Optional[str]:
        filename = paper.get("filename") or ""
        name = filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
        if name.lower().endswith(".pdf"):
            name = name[:-4]
        name = re.sub(r"[\s_]+", " ", name).strip()
        return name or None

    def _apply_record(self, paper: dict, record: dict, verified: bool, fill: bool = False) -> None:
        paper["verification_status"] = "verified" if verified else "unverified"
        if verified or fill:
            paper["doi"] = record.get("doi") or paper.get("doi")
            paper["year"] = record.get("year") or paper.get("year")
            paper["authors"] = record.get("authors") or paper.get("authors") or []
            paper["abstract"] = (
                metadata_sources.strip_xml(record.get("abstract")) or paper.get("abstract")
            )
            journal = record.get("journal") or record.get("container_title")
            paper["journal"] = (journal or "").strip() or paper.get("journal")
            paper["volume"] = (record.get("volume") or "").strip() or paper.get("volume")
            paper["issue"] = (record.get("issue") or "").strip() or paper.get("issue")
            paper["pages"] = (record.get("pages") or "").strip() or paper.get("pages")
            paper["publisher"] = (record.get("publisher") or "").strip() or paper.get("publisher")
            paper["url"] = (record.get("url") or "").strip() or paper.get("url")
            if record.get("title"):
                paper["title"] = record["title"]

    @staticmethod
    def _pick_best(query_title: str, candidates: List[dict]) -> Tuple[Optional[dict], float]:
        best, best_score = None, 0.0
        for record in candidates:
            record_title = record.get("title")
            if not record_title:
                continue
            score = fuzz.token_set_ratio(query_title.lower(), record_title.lower())
            if score > best_score:
                best, best_score = record, float(score)
        return best, best_score

    @staticmethod
    def _sanity_ok(
        paper: dict,
        first_page: Optional[str],
        record: dict,
        llm_hints: Optional[dict] = None,
    ) -> bool:
        """Author/year sanity gate before marking a title match as verified.

        If the PDF reveals a year (directly, or via the LLM's reading), a
        candidate more than two years off is rejected. If we know the paper's
        authors, a candidate sharing none of their family names is rejected.
        Unknown information is skipped, so a strong title match without
        author/year hints still verifies.
        """
        llm_hints = llm_hints or {}
        found_year = extract_year(first_page) if first_page else None
        if not found_year and llm_hints.get("year"):
            found_year = llm_hints["year"]
        if found_year and record.get("year") and abs(int(record["year"]) - int(found_year)) > 2:
            return False
        paper_authors = paper.get("authors") or []
        if not paper_authors:
            paper_authors = llm_hints.get("authors") or []
        record_authors = record.get("authors") or []
        if paper_authors and record_authors:
            paper_families = {LiteratureService._family(a) for a in paper_authors if a}
            record_families = {LiteratureService._family(a) for a in record_authors if a}
            paper_families.discard("")
            record_families.discard("")
            if paper_families and record_families and not (paper_families & record_families):
                return False
        return True

    @staticmethod
    def _family(name: str) -> str:
        if ", " in name:
            return name.split(",", 1)[0].strip().lower()
        parts = name.split()
        return parts[-1].lower() if parts else ""

    def _first_page_text(self, paper_id: str) -> Optional[str]:
        """Text of the first page of the stored PDF, or None."""
        row = self.db.execute(
            "SELECT file_path, file_type FROM documents WHERE id = ?",
            (paper_id,),
        ).fetchone()
        if not row or not row["file_path"]:
            return None
        path = Path(row["file_path"])
        if not path.is_file() or path.suffix.lower() != ".pdf":
            return None
        try:
            import fitz

            with fitz.open(str(path)) as doc:
                if not doc.page_count:
                    return None
                return doc[0].get_text() or None
        except Exception:
            return None

    def save_enrichment(self, paper: dict) -> None:
        self.db.execute(
            """UPDATE documents
               SET title = ?, year = ?, doi = ?, abstract = ?, authors = ?,
                   journal = ?, volume = ?, issue = ?, pages = ?, publisher = ?,
                   url = ?, verification_status = ?, apa_reference = ?,
                   extracted_doi = ?, metadata_candidates = ?
               WHERE id = ?""",
            (
                paper.get("title"),
                paper.get("year"),
                paper.get("doi"),
                paper.get("abstract"),
                json.dumps(paper.get("authors") or []),
                paper.get("journal"),
                paper.get("volume"),
                paper.get("issue"),
                paper.get("pages"),
                paper.get("publisher"),
                paper.get("url"),
                paper.get("verification_status"),
                self.apa_reference(paper),
                paper.get("extracted_doi"),
                json.dumps(paper.get("metadata_candidates") or []),
                paper["id"],
            ),
        )
        self.db.commit()

    @staticmethod
    def _strip_xml(text: str) -> str:
        if not text:
            return text
        return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", text)).strip()

    @staticmethod
    def apa_reference(paper: dict) -> str:
        authors = paper.get("authors") or []
        year = paper.get("year") or "n.d."
        title = (paper.get("title") or "Untitled").strip()
        if authors:
            shown = "; ".join(authors[:7])
            if len(authors) > 7:
                shown += " …"
            ref = f"{shown} ({year}). {title}."
        else:
            ref = f"{title} ({year})."

        journal = (paper.get("journal") or "").strip()
        if journal:
            ref += f" {journal}"
            volume = (paper.get("volume") or "").strip()
            issue = (paper.get("issue") or "").strip()
            if volume:
                ref += f", {volume}" + (f"({issue})" if issue else "")
            elif issue:
                ref += f" ({issue})"
        pages = (paper.get("pages") or "").strip()
        if pages:
            ref += f", {pages}"
        if journal or pages:
            ref += "."

        publisher = (paper.get("publisher") or "").strip()
        if publisher and not journal:
            ref += f" {publisher}."

        doi = (paper.get("doi") or "").strip().rstrip(".,")
        url = (paper.get("url") or "").strip().rstrip(".,")
        if doi:
            ref += f" https://doi.org/{doi}"
        elif url:
            ref += f" {url}"
        return ref

    # --- similarity edges ----------------------------------------------------

    def similarity_edges(self, user_id: str, papers: List[dict]) -> List[dict]:
        settings = SettingsService(self.db).get_settings(user_id).settings
        provider = resolve_embedding_provider(settings)
        texts = [
            self._embedding_text(p) for p in papers
        ]
        vectors = self._embed_batches(provider, texts)
        if not vectors or len(vectors) != len(papers):
            return []

        n = len(papers)
        sims: Dict[Tuple[str, str], float] = {}
        for i in range(n):
            for j in range(i + 1, n):
                sim = self._cosine(vectors[i], vectors[j])
                if sim > 0.0:
                    sims[(papers[i]["id"], papers[j]["id"])] = sim

        edges: Dict[Tuple[str, str], float] = {}
        for pair, sim in sims.items():
            if sim >= SIMILARITY_THRESHOLD:
                edges[pair] = round(sim, 3)
        for k in range(n):
            pid = papers[k]["id"]
            ranked = sorted(
                ((sim, a, b) for (a, b), sim in sims.items() if a == pid or b == pid),
                reverse=True,
            )
            for sim, a, b in ranked[: SIMILARITY_TOP_K]:
                pair = (a, b)
                edges[pair] = max(edges.get(pair, 0.0), round(sim, 3))

        return [
            {"source": a, "target": b, "weight": weight, "edge_type": "similarity"}
            for (a, b), weight in edges.items()
        ]

    @staticmethod
    def _embedding_text(paper: dict) -> str:
        title = (paper.get("title") or "").strip()
        abstract = (paper.get("abstract") or "").strip()
        if title and abstract:
            return f"{title}\n{abstract}"
        return title or abstract or "Untitled document"

    @staticmethod
    def _embed_batches(provider, texts: List[str], batch: int = 32) -> List[List[float]]:
        out: List[List[float]] = []
        for i in range(0, len(texts), batch):
            out.extend(provider.embed(texts[i : i + batch]))
        return out

    @staticmethod
    def _cosine(a: List[float], b: List[float]) -> float:
        if not a or not b or len(a) != len(b):
            return 0.0
        dot = sum(x * y for x, y in zip(a, b))
        na = math.sqrt(sum(x * x for x in a)) or 1.0
        nb = math.sqrt(sum(y * y for y in b)) or 1.0
        return dot / (na * nb)

    # --- citation edges ------------------------------------------------------

    def citation_edges(self, papers: List[dict]) -> Tuple[List[dict], List[dict], List[dict]]:
        by_id = {p["id"]: p for p in papers}
        edges: Dict[Tuple[str, str], float] = {}
        refs: List[dict] = []
        unmatched: List[dict] = []
        for paper in papers:
            for line in self._reference_lines(paper["id"]):
                match = self._match_title(line, by_id)
                if not match:
                    unmatched.append({"paper_id": paper["id"], "raw_ref": line})
                    continue
                matched_id, score = match
                if matched_id == paper["id"]:
                    continue
                pair = tuple(sorted((paper["id"], matched_id)))
                confidence = round(score / 100, 3)
                edges[pair] = max(edges.get(pair, 0.0), confidence)
                refs.append(
                    {
                        "paper_id": paper["id"],
                        "raw_ref": line,
                        "matched_paper_id": matched_id,
                        "confidence": confidence,
                    }
                )
        edge_rows = [
            {"source": a, "target": b, "weight": weight, "edge_type": "citation"}
            for (a, b), weight in edges.items()
        ]
        return edge_rows, refs, unmatched

    def save_references(self, refs: List[dict]) -> None:
        for ref in refs:
            self.db.execute(
                """INSERT INTO paper_references (id, paper_id, raw_ref, matched_paper_id, confidence)
                   VALUES (?, ?, ?, ?, ?)""",
                (
                    generate_id(),
                    ref["paper_id"],
                    ref["raw_ref"],
                    ref["matched_paper_id"],
                    ref["confidence"],
                ),
            )
        if refs:
            self.db.commit()

    def _reference_lines(self, paper_id: str) -> List[str]:
        rows = self.db.execute(
            """SELECT content, page_number FROM chunks
               WHERE document_id = ?
               ORDER BY page_number ASC, chunk_index ASC""",
            (paper_id,),
        ).fetchall()
        if not rows:
            return []
        text = "\n".join(r["content"] for r in rows)
        marker = re.search(r"(?i)\b(?:references|bibliography|works cited|literature cited)\b", text)
        if marker:
            section = text[marker.start() :]
        else:
            section = text[max(0, int(len(text) * 0.88)) :]

        blocks = re.split(r"\n\s*(?:\[\d+\]|\d{1,3}[.)])\s+", section)
        if len(blocks) == 1:
            blocks = section.splitlines()
        out = []
        for block in blocks:
            line = re.sub(r"\s+", " ", block).strip()
            if len(line) >= 12:
                out.append(line)
            if len(out) >= MAX_REFERENCE_LINES:
                break
        return out

    @staticmethod
    def _match_title(line: str, by_id: Dict[str, dict]) -> Optional[Tuple[str, float]]:
        best_id, best_score = None, 0.0
        for pid, paper in by_id.items():
            title = (paper.get("title") or "").strip()
            if len(title) < 10:
                continue
            score = fuzz.token_set_ratio(line.lower(), title.lower())
            if score > best_score:
                best_id, best_score = pid, float(score)
        if best_id is None or best_score < CITATION_THRESHOLD:
            return None
        return best_id, best_score

    # --- AI-judged edges -----------------------------------------------------

    def _ai_edge_pass(
        self,
        user_id: str,
        papers: List[dict],
        unmatched: List[dict],
        existing_edges: List[dict],
    ) -> Tuple[List[dict], List[dict]]:
        """LLM-judged similarity + citation edges that the embedding/fuzzy passes
        missed. Best-effort: returns empty when no model is configured or the
        calls fail, so the build never breaks on the LLM.
        """
        from app.features.literature.llm_service import LiteratureLLMService

        llm = LiteratureLLMService(self.db)
        if not llm.active_model(user_id):
            return [], []
        existing = {
            (e["source"], e["target"], e["edge_type"]) for e in (existing_edges or [])
        }
        edge_rows: List[dict] = []
        refs: List[dict] = []
        try:
            related = llm.related_pairs(user_id, papers)
            for edge in related:
                a, b = edge["source"], edge["target"]
                if (a, b, "similarity") in existing or (b, a, "similarity") in existing:
                    continue
                existing.add((a, b, "similarity"))
                edge_rows.append(edge)
        except Exception:
            pass
        try:
            matched = llm.match_reference_lines(user_id, unmatched, papers)
            seen = set()
            for ref in matched:
                pair = tuple(sorted((ref["paper_id"], ref["matched_paper_id"])))
                key = (pair[0], pair[1], "citation")
                if key in existing or (pair[1], pair[0], "citation") in existing or pair in seen:
                    continue
                seen.add(pair)
                refs.append(ref)
                edge_rows.append(
                    {
                        "source": pair[0],
                        "target": pair[1],
                        "weight": ref["confidence"],
                        "edge_type": "citation",
                    }
                )
        except Exception:
            pass
        return edge_rows, refs

    # --- clustering + layout -------------------------------------------------

    @staticmethod
    def assign_clusters(papers: List[dict], edges: List[dict]) -> Dict[str, str]:
        import networkx as nx
        from networkx.algorithms.community import greedy_modularity_communities

        graph = nx.Graph()
        graph.add_nodes_from(p["id"] for p in papers)
        graph.add_edges_from((e["source"], e["target"]) for e in edges)

        node_cluster: Dict[str, str] = {}
        for idx, community in enumerate(greedy_modularity_communities(graph), start=1):
            for node in community:
                node_cluster[node] = f"c{idx}"
        isolated = [p["id"] for p in papers if p["id"] not in node_cluster]
        for idx, node in enumerate(isolated, start=len(node_cluster) + 1):
            node_cluster[node] = f"c{idx}"
        return node_cluster

    # --- full build ----------------------------------------------------------

    def build_map(
        self,
        user_id: str,
        project_id: str,
        on_progress: Optional[Callable[[int, str], None]] = None,
    ) -> LiteratureMapResponse:
        def report(percent: int, stage: str) -> None:
            if on_progress:
                on_progress(percent, stage)

        report(5, "Loading papers")
        papers = self.papers(user_id, project_id)
        if not papers:
            return LiteratureMapResponse(nodes=[], edges=[], clusters=[])

        report(15, "Enriching metadata")
        total = len(papers)
        for idx, paper in enumerate(papers):
            try:
                self.enrich(paper, user_id=user_id)
                self.save_enrichment(paper)
            except Exception:
                continue
            report(15 + int((idx + 1) / total * 30), "Enriching metadata")

        report(45, "Embedding papers")
        similarity = self.similarity_edges(user_id, papers)

        report(60, "Matching citations")
        citations, refs, unmatched = self.citation_edges(papers)
        self.save_references(refs)

        report(66, "AI similarity & citations")
        ai_edges, ai_refs = self._ai_edge_pass(
            user_id, papers, unmatched, similarity + citations
        )
        self.save_references(ai_refs)

        report(72, "Clustering")
        edges = similarity + citations + ai_edges
        node_cluster = self.assign_clusters(papers, edges)

        report(80, "Computing layout")
        node_rows = [
            {
                "id": p["id"],
                "label": (p["title"] or "Untitled"),
                "type": "paper",
                "weight": 1.0,
                "cluster": node_cluster.get(p["id"]),
                "meta": {
                    "doc_id": p["id"],
                    "year": p.get("year"),
                    "authors": p.get("authors") or [],
                    "doi": p.get("doi"),
                    "abstract": p.get("abstract"),
                    "verification_status": p.get("verification_status"),
                    "apa_reference": p.get("apa_reference"),
                },
            }
            for p in papers
        ]
        node_rows = layout_nodes(node_rows, edges)
        nodes = [PaperNode(**n) for n in node_rows]
        edge_rows = [PaperEdge(**e) for e in edges]

        cluster_ids = sorted({c for c in node_cluster.values()})
        clusters = []
        for cid in cluster_ids:
            members = [n for n in nodes if n.cluster == cid]
            label = members[0].label if len(members) == 1 else f"Cluster {cid[1:]}"
            clusters.append(ClusterInfo(id=cid, label=label, summary="", size=len(members)))

        return LiteratureMapResponse(nodes=nodes, edges=edge_rows, clusters=clusters)

    # --- checkpoint persistence ----------------------------------------------

    def save_checkpoint(
        self,
        user_id: str,
        project_id: str,
        response: LiteratureMapResponse,
        fingerprint: str,
    ) -> None:
        self.db.execute(
            """INSERT INTO graph_history
                  (id, user_id, project_id, graph_json, fingerprint, prefs_key, map_type)
               VALUES (?, ?, ?, ?, ?, 'literature', 'literature')""",
            (generate_id(), user_id, project_id, response.model_dump_json(), fingerprint),
        )
        self.db.commit()

    def get_map(self, user_id: str, project_id: str) -> LiteratureMapResponse:
        try:
            row = self.db.execute(
                """SELECT graph_json FROM graph_history
                   WHERE user_id = ? AND project_id = ? AND map_type = 'literature'
                   ORDER BY created_at DESC, rowid DESC LIMIT 1""",
                (user_id, project_id),
            ).fetchone()
        except Exception:
            row = None
        if row is None:
            return LiteratureMapResponse(nodes=[], edges=[], clusters=[])
        data = json.loads(row["graph_json"])
        return LiteratureMapResponse(
            nodes=[PaperNode(**n) for n in data.get("nodes", [])],
            edges=[PaperEdge(**e) for e in data.get("edges", [])],
            clusters=[ClusterInfo(**c) for c in data.get("clusters", [])],
            generated_at=data.get("generated_at"),
        )
