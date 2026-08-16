import asyncio
import json
import sqlite3
from io import BytesIO
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from app.features.documents.repository import DocumentRepository
from app.features.literature import jobs
from app.features.literature.llm_service import LiteratureLLMService
from app.features.literature.service import LiteratureService
from app.features.search.schemas import SearchRequest
from app.features.search.service import SearchService

_SCHEMA = """
CREATE TABLE documents (
    id TEXT PRIMARY KEY, user_id TEXT, project_id TEXT, title TEXT, filename TEXT,
    file_path TEXT, file_type TEXT, author TEXT, year INTEGER, doi TEXT,
    abstract TEXT, verification_status TEXT, apa_reference TEXT, authors TEXT);
CREATE TABLE chunks (
    id TEXT PRIMARY KEY, document_id TEXT, chunk_index INTEGER, content TEXT,
    page_number INTEGER, char_start INTEGER, char_end INTEGER, token_count INTEGER,
    embedded_at TEXT);
CREATE TABLE user_settings (user_id TEXT PRIMARY KEY, settings TEXT NOT NULL DEFAULT '{}');
CREATE TABLE model_configs (
    id TEXT PRIMARY KEY, user_id TEXT, name TEXT, provider TEXT, model_id TEXT,
    is_active INTEGER DEFAULT 0, is_default INTEGER DEFAULT 0, config TEXT);
CREATE TABLE literature_entries (
    paper_id TEXT PRIMARY KEY, user_id TEXT, project_id TEXT, citation TEXT,
    research_objective TEXT, methodology TEXT, key_findings TEXT, limitations TEXT,
    relevance TEXT, apa_reference TEXT, auto_generated INTEGER DEFAULT 0,
    updated_at DATETIME DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE paper_references (
    id TEXT PRIMARY KEY, paper_id TEXT, raw_ref TEXT, matched_paper_id TEXT, confidence REAL);
"""


def _db():
    db = sqlite3.connect(":memory:", check_same_thread=False)
    db.row_factory = sqlite3.Row
    db.executescript(_SCHEMA)
    db.commit()
    return db


def _seed_paper(db, pid, project_id, title, author="", year=None, authors=None):
    db.execute(
        """INSERT INTO documents (id, user_id, project_id, title, author, year, authors, filename, file_path, file_type)
           VALUES (?, 'u1', ?, ?, ?, ?, ?, ?, ?, 'pdf')""",
        (pid, project_id, title, author, year, json.dumps(authors or []), f"{pid}.pdf", f"/tmp/{pid}.pdf"),
    )
    db.commit()


def _seed_chunk(db, cid, doc_id, content, index=0, page=None):
    db.execute(
        "INSERT INTO chunks (id, document_id, chunk_index, content, page_number) VALUES (?, ?, ?, ?, ?)",
        (cid, doc_id, index, content, page),
    )
    db.commit()


# --- entries ----------------------------------------------------------------


def test_ensure_entries_seeds_citation_and_apa():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey", author="Smith; Doe", year=2021, authors=["Smith, John", "Doe, Jane"])
    service = LiteratureService(db)

    assert service.ensure_entries("u1", "proj1") == 1
    entry = service.get_entry("p1", "u1")
    assert entry["citation"] == "Smith, 2021"
    assert entry["paper_id"] == "p1"
    assert entry["apa_reference"].startswith("Smith, John; Doe, Jane (2021). Deep Learning Survey.")
    assert entry["auto_generated"] is True


def test_auto_upsert_never_overwrites_user_edits():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey", author="Smith", year=2021)
    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")

    edited = service.upsert_entry(
        "p1", "u1", "proj1", {"research_objective": "My own rewritten objective"}, auto=False
    )
    assert edited["research_objective"] == "My own rewritten objective"
    assert edited["auto_generated"] is False

    again = service.upsert_entry(
        "p1", "u1", "proj1", {"research_objective": "AI overwrite attempt", "methodology": "AI method"}, auto=True
    )
    assert again["research_objective"] == "My own rewritten objective"
    assert again["methodology"] == "AI method"
    assert again["auto_generated"] is False


def test_user_update_can_clear_a_field():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey")
    service = LiteratureService(db)
    service.upsert_entry("p1", "u1", "proj1", {"key_findings": "Something"}, auto=False)

    cleared = service.upsert_entry("p1", "u1", "proj1", {"key_findings": ""}, auto=False)
    assert cleared["key_findings"] == ""


def test_auto_citation_formats():
    assert LiteratureService.auto_citation({"authors": ["Smith, Jane"], "year": 2020}) == "Smith, 2020"
    assert LiteratureService.auto_citation({"authors": ["Ada Lovelace"], "year": None}) == "Lovelace, n.d."
    assert LiteratureService.auto_citation({"author": "Jones; Green", "year": 1999, "authors": []}) == "Jones, 1999"
    assert LiteratureService.auto_citation({"title": "No Authors Here", "year": 2001, "authors": []}) == "No Authors Here, 2001"


# --- export -----------------------------------------------------------------


def test_export_rows_falls_back_to_auto_values():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey", author="Smith", year=2021, authors=["Smith, Jane"])
    _seed_paper(db, "p2", "proj1", "Another Study", author="Jones", year=2020, authors=["Jones, Bob"])
    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")
    service.upsert_entry("p2", "u1", "proj1", {"research_objective": "Research Q for p2"}, auto=False)

    rows = service.export_rows("u1", "proj1")
    by_title = {r["title"]: r for r in rows}
    assert set(by_title) == {"Deep Learning Survey", "Another Study"}
    assert by_title["Deep Learning Survey"]["citation"] == "Smith, 2021"
    assert by_title["Another Study"]["research_objective"] == "Research Q for p2"
    assert "Deep Learning Survey" in by_title["Deep Learning Survey"]["apa_reference"]
    assert "Another Study" in by_title["Another Study"]["apa_reference"]


def test_export_workbook_structure():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey", author="Smith", year=2021)
    _seed_paper(db, "p2", "proj1", "Another Study", author="Jones", year=2020)
    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")

    data = service.export_workbook("u1", "proj1")
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["Literature Mapping"]

    headers = [sheet.cell(row=1, column=c).value for c in range(1, 9)]
    assert headers[0] == "Paper Title"
    assert headers[1] == "Citation (Author, Year)"
    assert headers[-1] == "APA Reference"
    assert sheet.max_row == 3  # header + 2 papers
    citations = {sheet.cell(row=r, column=2).value for r in range(2, 4)}
    assert {"Smith, 2021", "Jones, 2020"} <= citations


# --- LLM entry generation ---------------------------------------------------


def test_summarize_papers_fills_entry_via_fallback_without_model():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey")
    db.execute(
        "UPDATE documents SET abstract = ? WHERE id = 'p1'",
        ("We introduce a transformer architecture.",),
    )
    db.commit()
    _seed_chunk(db, "c1", "p1", "We introduce a transformer architecture.", index=0, page=1)
    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")
    llm = LiteratureLLMService(db)

    with patch.object(LiteratureLLMService, "active_model", return_value=None):
        assert llm.summarize_papers("u1", "proj1") == 1
    entry = service.get_entry("p1", "u1")
    assert "transformer architecture" in entry["research_objective"].lower()
    assert "LLM unavailable" in entry["limitations"]


def test_summarize_paper_regenerates_with_llm_output():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey", year=2021)
    _seed_chunk(db, "c1", "p1", "We study generalization of deep networks.", index=0, page=1)
    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")
    llm = LiteratureLLMService(db)

    async def fake_call(self, user_id, system, prompt):
        assert "generalization" in prompt
        return json.dumps(
            {
                "research_objective": "Examine generalization.",
                "methodology": "Experiments on CNNs.",
                "key_findings": "Dropout helps.",
                "limitations": "Small sample.",
                "relevance": "Supports our method.",
            }
        )

    with patch.object(LiteratureLLMService, "active_model", return_value={"id": "m1"}), patch.object(
        LiteratureLLMService, "_call", new=fake_call
    ):
        entry = asyncio.run(llm.summarize_paper("u1", "proj1", "p1"))

    assert entry["research_objective"] == "Examine generalization."
    assert entry["relevance"] == "Supports our method."
    assert entry["auto_generated"] is True


def test_summarize_paper_raises_without_model():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey")
    llm = LiteratureLLMService(db)
    with patch.object(LiteratureLLMService, "active_model", return_value=None):
        with pytest.raises(ValueError):
            asyncio.run(llm.summarize_paper("u1", "proj1", "p1"))


# --- jobs -------------------------------------------------------------------


def test_start_build_sets_rebuild_pending_for_running_job():
    jobs._jobs["fake"] = {
        "id": "fake",
        "user_id": "u1",
        "project_id": "proj1",
        "status": "running",
        "progress": 50,
        "stage": "Enriching metadata",
        "error": None,
    }
    try:
        job_id = jobs.start_build(lambda: object(), "u1", "proj1")
        assert job_id == "fake"
        with jobs._jobs_lock:
            assert jobs._jobs["fake"]["rebuild_pending"] is True
        assert "rebuild_pending" not in jobs.get_job("fake")
    finally:
        with jobs._jobs_lock:
            jobs._jobs.pop("fake", None)


# --- editable metadata ------------------------------------------------------


def test_update_metadata_regenerates_apa_and_entry():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey", author="Smith", year=2021, authors=["Smith, John"])
    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")
    assert service.get_entry("p1", "u1")["citation"] == "Smith, 2021"

    meta = service.update_metadata(
        "p1",
        "u1",
        "proj1",
        {"authors": ["Zhang, Wei", "Kim, Soo"], "year": 2023, "doi": "10.1000/xyz"},
    )
    assert meta["authors"] == ["Zhang, Wei", "Kim, Soo"]
    assert meta["year"] == 2023
    assert meta["doi"] == "10.1000/xyz"
    assert meta["apa_reference"].startswith("Zhang, Wei; Kim, Soo (2023).")
    assert "https://doi.org/10.1000/xyz" in meta["apa_reference"]
    assert meta["metadata_user_edited"] is True

    entry = service.get_entry("p1", "u1")
    assert entry["citation"] == "Zhang, 2023"
    assert entry["apa_reference"] == meta["apa_reference"]


def test_metadata_update_preserves_user_overridden_apa():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey", author="Smith", year=2021)
    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")
    service.upsert_entry("p1", "u1", "proj1", {"apa_reference": "My custom APA string"}, auto=False)

    meta = service.update_metadata("p1", "u1", "proj1", {"year": 2024})
    entry = service.get_entry("p1", "u1")
    assert entry["apa_reference"] == "My custom APA string"  # user override kept
    assert entry["citation"] == "Smith, 2024"  # derived citation refreshed
    assert "2024" in meta["apa_reference"]


def test_enrich_skips_user_edited_metadata():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey", author="Smith", year=2021)
    service = LiteratureService(db)
    service.update_metadata("p1", "u1", "proj1", {"authors": ["Manual, Author"], "year": 1999})

    paper = next(p for p in service.papers("u1", "proj1") if p["id"] == "p1")
    with patch("app.features.literature.metadata.requests.get") as mock_get:
        enriched = service.enrich(paper)

    assert enriched["year"] == 1999
    assert enriched["authors"] == ["Manual, Author"]
    assert enriched["verification_status"] is None
    mock_get.assert_not_called()


# --- keyword search ---------------------------------------------------------


def test_keyword_search_finds_terms_and_merges_before_semantic():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey")
    _seed_paper(db, "p2", "proj1", "Gardening Guide")
    _seed_chunk(db, "c1", "p1", "Dropout regularization improves generalization in deep networks.", index=0, page=3)
    _seed_chunk(db, "c2", "p2", "Plant the seeds in March.", index=0, page=1)

    service = SearchService(db)
    req = SearchRequest(query="dropout regularization generalization", top_k=5, project_id="proj1")
    with patch.object(service, "_semantic_sync", return_value=[]):
        response = asyncio.run(service.search(req, "u1"))

    assert len(response.results) == 1
    assert response.results[0].chunk_id == "c1"
    assert response.results[0].source == "keyword"
    assert response.results[0].document_id == "p1"
    assert response.results[0].page_number == 3
    assert response.results[0].score > 0


def test_keyword_search_merges_semantic_first_and_dedupes():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey")
    _seed_chunk(db, "c1", "p1", "Dropout regularization improves generalization.", index=0, page=1)
    service = SearchService(db)
    semantic = [
        {
            "chunk_id": "c1",
            "content": "Dropout regularization improves generalization.",
            "score": 0.9,
            "document_id": "p1",
            "document_title": "Deep Learning Survey",
            "page_number": 1,
            "source": "semantic",
        }
    ]
    req = SearchRequest(query="dropout", top_k=5, project_id="proj1")
    with patch.object(service, "_semantic_sync", return_value=semantic):
        response = asyncio.run(service.search(req, "u1"))

    assert len(response.results) == 1
    assert response.results[0].source == "semantic"


# --- delete cascade ---------------------------------------------------------


def test_delete_removes_literature_data_for_the_paper():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey")
    _seed_paper(db, "p2", "proj1", "Another Study")
    _seed_chunk(db, "c1", "p1", "Body text", index=0, page=1)
    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")
    service.save_references(
        [
            {"paper_id": "p1", "raw_ref": "Another Study 2020", "matched_paper_id": "p2", "confidence": 0.9},
            {"paper_id": "p2", "raw_ref": "Deep Learning Survey 2021", "matched_paper_id": "p1", "confidence": 0.9},
        ]
    )

    DocumentRepository(db).delete("p1", "u1")

    assert service.get_entry("p1", "u1") is None
    assert service.get_entry("p2", "u1") is not None
    remaining = db.execute("SELECT * FROM paper_references").fetchall()
    assert len(remaining) == 0  # both directions removed
    chunks = db.execute("SELECT * FROM chunks").fetchall()
    assert len(chunks) == 0
    docs = db.execute("SELECT id FROM documents").fetchall()
    assert [d["id"] for d in docs] == ["p2"]


# --- LLM-unavailable fallback ------------------------------------------------


def test_summarize_papers_fills_entry_from_abstract_when_llm_down():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Farm Budgets", authors=["Kibirige, D."], year=2014)
    db.execute(
        """UPDATE documents SET abstract = ? WHERE id = 'p1'""",
        (
            "This study estimates the budgets of small-scale farms in the Eastern Cape. "
            "A survey of 120 households was conducted. Results show higher margins "
            "for commercial farms.",
        ),
    )
    db.execute(
        "INSERT INTO model_configs (id, user_id, name, provider, model_id, is_active) "
        "VALUES ('m1', 'u1', 'llama', 'ollama', 'llama', 1)"
    )
    db.commit()

    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")
    llm = LiteratureLLMService(db)

    with patch.object(llm, "_call", side_effect=RuntimeError("model down")):
        llm.summarize_papers("u1", "proj1")

    entry = service.get_entry("p1", "u1")
    assert "Eastern Cape" in entry["research_objective"]
    assert "survey" in entry["methodology"].lower()
    assert "margins" in entry["key_findings"].lower()
    assert "LLM unavailable" in entry["limitations"]
    assert entry["auto_generated"] is True


def test_summarize_papers_replaces_fallback_when_llm_returns():
    db = _db()
    _seed_paper(db, "p1", "proj1", "Deep Learning Survey", authors=["Smith, J."], year=2021)
    db.execute(
        "UPDATE documents SET abstract = ? WHERE id = 'p1'",
        ("An approach for improving deep learning accuracy.",),
    )
    db.execute(
        "INSERT INTO model_configs (id, user_id, name, provider, model_id, is_active) "
        "VALUES ('m1', 'u1', 'llama', 'ollama', 'llama', 1)"
    )
    db.commit()

    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")
    llm = LiteratureLLMService(db)

    with patch.object(llm, "_call", side_effect=RuntimeError("model down")):
        llm.summarize_papers("u1", "proj1")
    entry = service.get_entry("p1", "u1")
    assert "LLM unavailable" in entry["limitations"]

    good = json.dumps(
        {
            "research_objective": "Objective A",
            "methodology": "Method B",
            "key_findings": "Findings C",
            "limitations": "Limits D",
            "relevance": "Relevance E",
        }
    )
    with patch.object(llm, "_call", return_value=good):
        llm.summarize_papers("u1", "proj1")

    entry = service.get_entry("p1", "u1")
    assert entry["research_objective"] == "Objective A"
    assert entry["methodology"] == "Method B"
    assert "LLM unavailable" not in entry["limitations"]


# --- reference metadata fields ------------------------------------------------


def test_apa_reference_includes_journal_volume_issue_pages_publisher_url():
    paper = {
        "title": "A Study of Farming",
        "authors": ["Kibirige, D."],
        "year": 2014,
        "journal": "Journal of Agriculture",
        "volume": "12",
        "issue": "3",
        "pages": "44-60",
        "publisher": "Agri Press",
        "url": "https://example.org/paper",
    }
    ref = LiteratureService.apa_reference(paper)
    assert "Kibirige, D. (2014). A Study of Farming." in ref
    assert "Journal of Agriculture, 12(3), 44-60." in ref
    assert "https://example.org/paper" in ref

    no_journal = LiteratureService.apa_reference({**paper, "journal": None, "volume": None, "issue": None, "pages": None})
    assert "Agri Press." in no_journal

    with_doi = LiteratureService.apa_reference({**paper, "doi": "10.1000/xyz"})
    assert "https://doi.org/10.1000/xyz" in with_doi
    assert "https://example.org/paper" not in with_doi


def test_update_metadata_persists_reference_fields_and_refreshes_apa():
    db = _db()
    _seed_paper(db, "p1", "proj1", "A Study of Farming", authors=["Kibirige, D."], year=2014)
    service = LiteratureService(db)
    service.ensure_entries("u1", "proj1")

    meta = service.update_metadata(
        "p1", "u1", "proj1",
        {
            "journal": "Journal of Agriculture",
            "volume": "12",
            "issue": "3",
            "pages": "44-60",
            "publisher": "Agri Press",
            "url": "https://example.org/paper",
        },
    )
    assert meta is not None
    assert meta["journal"] == "Journal of Agriculture"
    assert meta["volume"] == "12"
    assert meta["issue"] == "3"
    assert meta["pages"] == "44-60"
    assert meta["publisher"] == "Agri Press"
    assert meta["url"] == "https://example.org/paper"
    assert "Journal of Agriculture, 12(3), 44-60." in meta["apa_reference"]

    entry = service.get_entry("p1", "u1")
    assert "Journal of Agriculture, 12(3), 44-60." in entry["apa_reference"]
